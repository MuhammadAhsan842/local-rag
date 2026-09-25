"""
Render a tailored, ATS-clean CV (gap #1).

Research: two-column / graphic PDFs parse at ~71% vs ~97% for single-column —
and if the ATS can't extract your content, your match score is low regardless
of fit. So the tailored CV we submit must be *single-column, no tables, no text
boxes, real sentences*. This module renders exactly that from verified facts +
tailored bullets, in three ATS-safe formats:

  .txt   — plain text, the most parser-proof of all
  .html  — single-column, print-to-PDF clean (no columns/tables/graphics)
  .docx  — only if python-docx is installed (optional)

Everything rendered is grounded: bullets come from tailor.py (faithfulness-gated),
skills/contact from your profile. Nothing is invented here.
"""
from __future__ import annotations

import html
from pathlib import Path


def _sections(contact: dict, summary: str, bullets: list[dict], facts: list[dict],
              skills: list[str], job_title: str, company: str) -> dict:
    fact_map = {f["id"]: f for f in facts}
    exp = []
    for b in bullets:
        src = fact_map.get(b["fact_id"], {})
        exp.append({"text": b["text"], "source": src.get("source", "")})
    return {
        "name": contact.get("name", ""),
        "contactline": " | ".join(x for x in [
            contact.get("email", ""), contact.get("phone", ""),
            contact.get("location", ""), contact.get("linkedin", ""),
            contact.get("github", "")] if x),
        "target": f"{job_title} — {company}".strip(" —"),
        "summary": summary,
        "experience": exp,
        "skills": skills,
    }


def render_text(s: dict) -> str:
    L = [s["name"], s["contactline"], ""]
    if s["target"]:
        L += [f"TARGET ROLE: {s['target']}", ""]
    if s["summary"]:
        L += ["SUMMARY", s["summary"], ""]
    L += ["RELEVANT EXPERIENCE"]
    for e in s["experience"]:
        tag = f"  ({e['source']})" if e["source"] else ""
        L.append(f"- {e['text']}{tag}")
    if s["skills"]:
        L += ["", "SKILLS", ", ".join(s["skills"])]
    return "\n".join(L).strip() + "\n"


def render_html(s: dict) -> str:
    def esc(x): return html.escape(str(x))
    exp = "\n".join(
        f"      <li>{esc(e['text'])}"
        + (f" <span class='src'>({esc(e['source'])})</span>" if e["source"] else "")
        + "</li>"
        for e in s["experience"])
    skills = f"<p>{esc(', '.join(s['skills']))}</p>" if s["skills"] else ""
    summary = f"<h2>Summary</h2><p>{esc(s['summary'])}</p>" if s["summary"] else ""
    target = f"<p class='target'>Target: {esc(s['target'])}</p>" if s["target"] else ""
    # deliberately single-column, no tables/floats/graphics — ATS-parser safe
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<title>{esc(s['name'])} — CV</title>
<style>
  body {{ font: 12pt/1.45 Georgia, 'Times New Roman', serif; color:#111;
          max-width: 720px; margin: 32px auto; padding: 0 16px; }}
  h1 {{ font-size: 20pt; margin: 0 0 2px; }}
  h2 {{ font-size: 12.5pt; text-transform: uppercase; letter-spacing:.5px;
        border-bottom: 1px solid #999; padding-bottom: 2px; margin: 18px 0 6px; }}
  .contact, .target {{ color:#333; font-size: 10.5pt; margin: 2px 0; }}
  ul {{ margin: 4px 0; padding-left: 20px; }} li {{ margin: 4px 0; }}
  .src {{ color:#777; font-size: 9.5pt; }}
  @media print {{ body {{ margin: 0; }} }}
</style></head><body>
  <h1>{esc(s['name'])}</h1>
  <p class="contact">{esc(s['contactline'])}</p>
  {target}
  {summary}
  <h2>Relevant Experience</h2>
  <ul>
{exp}
  </ul>
  <h2>Skills</h2>
  {skills}
</body></html>
"""


def render(contact: dict, tailored: dict, facts: list[dict], skills: list[str],
           job_title: str = "", company: str = "", summary: str = "",
           out_base: str = "tailored_cv") -> dict:
    """
    Write tailored CV files. Returns {"txt": path, "html": path, "docx": path?}.
    Always writes .txt + .html (ATS-safe, no deps). Writes .docx if python-docx
    is available.
    """
    s = _sections(contact, summary, tailored.get("bullets", []), facts, skills,
                  job_title, company)
    parent = Path(out_base).parent
    if str(parent) not in (".", ""):
        parent.mkdir(parents=True, exist_ok=True)
    out = {}
    txt_p = f"{out_base}.txt"
    Path(txt_p).write_text(render_text(s))
    out["txt"] = txt_p
    html_p = f"{out_base}.html"
    Path(html_p).write_text(render_html(s))
    out["html"] = html_p
    try:
        from docx import Document  # optional
        doc = Document()
        doc.add_heading(s["name"], level=0)
        doc.add_paragraph(s["contactline"])
        if s["summary"]:
            doc.add_heading("Summary", level=1)
            doc.add_paragraph(s["summary"])
        doc.add_heading("Relevant Experience", level=1)
        for e in s["experience"]:
            doc.add_paragraph(e["text"], style="List Bullet")
        if s["skills"]:
            doc.add_heading("Skills", level=1)
            doc.add_paragraph(", ".join(s["skills"]))
        docx_p = f"{out_base}.docx"
        doc.save(docx_p)
        out["docx"] = docx_p
    except Exception:  # noqa: BLE001 - python-docx not installed; txt+html suffice
        pass
    return out
